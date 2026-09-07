"""End-to-end HTTP journey through staged offline production inference."""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document

import pytest
from fastapi.testclient import TestClient

from mudidi.web.app import _preset_asset_links, _read_preset_text, create_app
from mudidi.config.yaml_config import InferenceConfig
from mudidi.web.credentials import CredentialVault
from mudidi.web.models import Provider
from mudidi.web.runs import RunStatus

def _preview(
    client: TestClient,
    tmp_path: Path,
    **overrides: str,
) -> str:
    data = {
        "output_directory": str(tmp_path / "output"),
        "pipeline": "complete",
        "provider": "anthropic",
        "model": "anthropic/claude-sonnet-5",
        "reasoning": "low",
        "agentic": "true",
        "verify_stage1": "true",
        "verify_stage2": "true",
        "parse_rules_pages": "1",
        "dictionary_pages": "1",
    }
    data.update(overrides)
    response = client.post(
        "/runs/preview",
        data=data,
        files={
            "dictionary_pdf": ("dictionary.pdf", _pdf_bytes(), "application/pdf")
        },
    )
    assert response.status_code == 200
    match = re.search(r'action="/runs/([^/]+)/start"', response.text)
    assert match is not None
    return match.group(1)


def _pdf_bytes(page_count: int = 1) -> bytes:
    import fitz

    document = fitz.open()
    for _ in range(page_count):
        document.new_page()
    try:
        return document.tobytes()
    finally:
        document.close()


def test_review_page_shows_page_ranges_and_each_stage_model(tmp_path: Path) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    response = TestClient(app).post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "output"),
            "pipeline": "complete",
            "dictionary_pages": "10-12",
            "parse_rules_pages": "10,12",
            "provider": "anthropic",
            "stage1_model": "gemini/gemini-3.1-pro-preview",
            "stage2_pass1_model": "anthropic/claude-opus-4-6",
            "stage2_pass2_model": "openai/gpt-5.4",
            "reasoning": "low",
            "agentic": "false",
        },
        files={
            "dictionary_pdf": (
                "dictionary.pdf",
                _pdf_bytes(12),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 200
    for heading in ("Input", "Pipeline", "Model", "Agentic"):
        assert f">{heading}<" in response.text
    assert 'class="review-actions panel"' in response.text
    assert "Start run" in response.text
    assert "Save these non-secret settings as a preset" in response.text
    assert "Dictionary Pages" in response.text
    assert "10-12" in response.text
    assert "Parse Rule Pages" in response.text
    assert "10, 12" in response.text
    assert "Stage 1 Model" in response.text
    assert "gemini/gemini-3.1-pro-preview" in response.text
    assert "Stage 2 Pass 1 Model" in response.text
    assert "anthropic/claude-opus-4-6" in response.text
    assert "Stage 2 Pass 2 Model" in response.text
    assert "openai/gpt-5.4" in response.text

@pytest.mark.parametrize(
    ("pipeline", "stage1_instructions", "stage2_instructions"),
    [
        ("transcription", "Mark uncertain letters.", None),
        ("structure", None, "Use the custom nt marker."),
        (
            "complete",
            "Mark uncertain letters.",
            "Use the custom nt marker.",
        ),
    ],
)
def test_preview_review_summarizes_enabled_instruction_metadata(
    tmp_path: Path,
    pipeline: str,
    stage1_instructions: str | None,
    stage2_instructions: str | None,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    data = {
        "output_directory": str(tmp_path / "output"),
        "pipeline": pipeline,
        "provider": "anthropic",
        "model": "anthropic/claude-sonnet-5",
        "reasoning": "low",
        "dictionary_pages": "1",
    }
    if stage1_instructions is not None:
        data["stage1_additional_instructions"] = stage1_instructions
    if stage2_instructions is not None:
        data["stage2_additional_instructions"] = stage2_instructions

    response = client.post(
        "/runs/preview",
        data=data,
        files={"dictionary_pdf": ("dictionary.pdf", _pdf_bytes(), "application/pdf")},
    )

    assert response.status_code == 200
    stage1_match = re.search(
        r"<dt>Stage 1 Instructions</dt>\s*<dd>(.*?)</dd>",
        response.text,
        re.S,
    )
    stage2_match = re.search(
        r"<dt>Stage 2 Instructions</dt>\s*<dd>(.*?)</dd>",
        response.text,
        re.S,
    )
    assert stage1_match is not None
    assert stage2_match is not None
    assert ("Source: Typed" in stage1_match.group(1)) == (
        stage1_instructions is not None
    )
    assert ("Source: Typed" in stage2_match.group(1)) == (
        stage2_instructions is not None
    )
    assert all(
        text not in response.text
        for text in (stage1_instructions, stage2_instructions)
        if text
    )


def test_prepared_review_preserves_instruction_metadata_summary(
    tmp_path: Path,
) -> None:
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=CredentialVault(environ={}),
        offline_inference=True,
    )
    client = TestClient(app)
    run_id = _preview(
        client,
        tmp_path,
        stage1_additional_instructions="Keep uncertain letters marked.",
        stage2_additional_instructions="Use the custom nt marker.",
    )

    recovered = client.get(f"/runs/{run_id}/review")

    assert recovered.status_code == 200
    assert "Stage 1 Instructions" in recovered.text
    assert "Stage 2 Instructions" in recovered.text
    assert "Source: Typed" in recovered.text
    assert "Keep uncertain letters marked." not in recovered.text
    assert "Use the custom nt marker." not in recovered.text
@pytest.mark.parametrize(
    ("inherited_stage1", "inherited_stage2", "pipeline"),
    [
        ("Saved Stage 1 instructions", None, "complete"),
        (None, "Saved Stage 2 instructions", "complete"),
        (
            "Saved Stage 1 instructions",
            "Saved Stage 2 instructions",
            "complete",
        ),
        (
            "Saved Stage 1 instructions",
            "Saved Stage 2 instructions",
            "transcription",
        ),
        (
            "Saved Stage 1 instructions",
            "Saved Stage 2 instructions",
            "structure",
        ),
    ],
)
def test_reused_preset_cleared_instructions_match_recovered_review(
    tmp_path: Path,
    inherited_stage1: str | None,
    inherited_stage2: str | None,
    pipeline: str,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    source_run_id = _preview(
        client,
        tmp_path,
        stage1_additional_instructions=inherited_stage1,
        stage2_additional_instructions=inherited_stage2,
    )
    saved = client.post(
        f"/runs/{source_run_id}/presets",
        data={"name": "Inherited instructions"},
        follow_redirects=False,
    )
    assert saved.status_code == 303
    preset = app.state.run_store.list_presets()[0]

    immediate = client.post(
        "/runs/preview",
        data={
            "preset_id": preset.preset_id,
            "output_directory": str(tmp_path / f"{pipeline}-output"),
            "pipeline": pipeline,
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "agentic": "false",
            "dictionary_pages": "1",
            "parse_rules_pages": "1",
            "stage1_additional_instructions": "",
            "stage2_additional_instructions": "",
        },
    )

    assert immediate.status_code == 200
    recovered_run = next(
        run for run in app.state.run_store.list_runs() if run.run_id != source_run_id
    )
    recovered = client.get(f"/runs/{recovered_run.run_id}/review")
    assert recovered.status_code == 200

    def instruction_rows(html: str) -> dict[str, str]:
        rows: dict[str, str] = {}
        for stage in ("Stage 1", "Stage 2"):
            match = re.search(
                rf"<dt>{stage} Instructions</dt>\s*<dd>(.*?)</dd>",
                html,
                re.S,
            )
            assert match is not None
            rows[stage] = match.group(1)
        return rows

    immediate_rows = instruction_rows(immediate.text)
    recovered_rows = instruction_rows(recovered.text)
    assert immediate_rows == recovered_rows
    assert all("Source: None" in row for row in immediate_rows.values())
    assert all(
        text not in immediate.text and text not in recovered.text
        for text in (inherited_stage1, inherited_stage2)
        if text
    )





def test_prepared_review_recovery_route_is_read_only_and_safe(
    tmp_path: Path,
) -> None:
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=CredentialVault(environ={}),
        offline_inference=True,
    )
    client = TestClient(app)
    run_id = _preview(client, tmp_path)
    prepared_status = app.state.run_store.get_run(run_id).status
    config_mtime = app.state.job_controller.config_path(run_id).stat().st_mtime_ns

    recovered = client.get(f"/runs/{run_id}/review")

    assert recovered.status_code == 200
    assert "Review your run" in recovered.text
    assert 'class="review-actions panel"' in recovered.text
    assert app.state.run_store.get_run(run_id).status is prepared_status
    assert app.state.job_controller.config_path(run_id).stat().st_mtime_ns == config_mtime

    blocked = client.post(f"/runs/{run_id}/start")
    assert blocked.status_code == 409
    assert app.state.run_store.get_run(run_id).status is RunStatus.CREDENTIALS_REQUIRED
    recovered_after_block = client.get(f"/runs/{run_id}/review")
    assert recovered_after_block.status_code == 200
    assert f'action="/runs/{run_id}/start"' in recovered_after_block.text
    assert "Start run" in recovered_after_block.text
    assert f'action="/runs/{run_id}/resume"' not in recovered_after_block.text


def test_recovered_review_labels_uploaded_guide_as_direct_use(
    tmp_path: Path,
) -> None:
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=CredentialVault(environ={}),
        offline_inference=True,
    )
    client = TestClient(app)
    response = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "output"),
            "pipeline": "complete",
            "dictionary_pages": "1",
            "parse_rules_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "agentic": "false",
        },
        files={
            "dictionary_pdf": ("dictionary.pdf", _pdf_bytes(), "application/pdf"),
            "existing_mdf_guide_file": (
                "uploaded-guide.json",
                b'{"markers": [{"marker": "lx", "description": "Headword"}]}',
                "application/json",
            ),
        },
    )

    assert response.status_code == 200
    run_id = re.search(r'action="/runs/([^/]+)/start"', response.text)
    assert run_id is not None
    recovered = client.get(f"/runs/{run_id.group(1)}/review")

    assert recovered.status_code == 200
    assert "Uploaded guide used directly" in recovered.text
    assert "Human approval required" not in recovered.text


def test_credential_blocked_resume_preserves_phase_and_approval_provenance(
    tmp_path: Path,
) -> None:
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=CredentialVault(environ={}),
        offline_inference=True,
    )
    client = TestClient(app)

    parse_run_id = _preview(client, tmp_path)
    app.state.run_store.transition(parse_run_id, RunStatus.QUEUED)
    app.state.run_store.transition(parse_run_id, RunStatus.DISCOVERING_PARSE_RULES)
    app.state.run_store.interrupt(parse_run_id)

    parse_blocked = client.post(f"/runs/{parse_run_id}/resume")

    assert parse_blocked.status_code == 409
    assert (
        f'data-continue-action="/runs/{parse_run_id}/resume"'
        in parse_blocked.text
    )
    parse_run = app.state.run_store.get_run(parse_run_id)
    assert parse_run.status is RunStatus.CREDENTIALS_REQUIRED
    assert parse_run.resume_phase == "parse_rule_review"
    parse_review = client.get(f"/runs/{parse_run_id}/review")
    assert parse_review.status_code == 200
    assert f'action="/runs/{parse_run_id}/resume"' in parse_review.text
    assert "Resume run" in parse_review.text

    pass2_run_id = _preview(client, tmp_path)
    app.state.run_store.transition(pass2_run_id, RunStatus.QUEUED)
    app.state.run_store.transition(
        pass2_run_id,
        RunStatus.AWAITING_PARSE_RULES_REVIEW,
    )
    app.state.run_store.authorize_pass2(
        pass2_run_id,
        review_id="review-provenance",
        approval_digest="a" * 64,
    )
    app.state.run_store.interrupt(pass2_run_id)

    pass2_blocked = client.post(f"/runs/{pass2_run_id}/resume")

    assert pass2_blocked.status_code == 409
    assert (
        f'data-continue-action="/runs/{pass2_run_id}/resume"'
        in pass2_blocked.text
    )
    pass2_review = client.get(f"/runs/{pass2_run_id}/review")
    assert pass2_review.status_code == 200
    assert f'action="/runs/{pass2_run_id}/resume"' in pass2_review.text
    assert "Resume run" in pass2_review.text
    pass2_run = app.state.run_store.get_run(pass2_run_id)
    assert pass2_run.status is RunStatus.CREDENTIALS_REQUIRED
    assert pass2_run.resume_phase == "stage2_pass2"
    assert pass2_run.review_id == "review-provenance"
    assert pass2_run.approval_digest == "a" * 64

    app.state.credential_vault.set_temporary(Provider.ANTHROPIC, "sk-ant-resume")
    continued = client.post(f"/runs/{parse_run_id}/resume", follow_redirects=False)

    assert continued.status_code == 303
    assert app.state.run_store.get_run(
        parse_run_id
    ).status is RunStatus.AWAITING_PARSE_RULES_REVIEW


def test_initial_credential_block_uses_start_continuation(
    tmp_path: Path,
) -> None:
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=CredentialVault(environ={}),
        offline_inference=True,
    )
    client = TestClient(app)
    run_id = _preview(client, tmp_path)

    blocked = client.post(f"/runs/{run_id}/start")

    assert blocked.status_code == 409
    assert f'data-continue-action="/runs/{run_id}/start"' in blocked.text

def test_review_recovery_returns_safe_404_for_unknown_or_unprepared_runs(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data")
    client = TestClient(app)
    app.state.run_store.create_run("run-unprepared", provider="anthropic")

    assert client.get("/runs/not-a-real-run/review").status_code == 404
    assert client.get("/runs/run-unprepared/review").status_code == 404


def test_review_start_pause_approve_and_complete_offline_journey(
    tmp_path: Path,
) -> None:
    vault = CredentialVault(environ={})
    vault.set_temporary(Provider.ANTHROPIC, "sk-ant-offline")
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=vault,
        offline_inference=True,
    )
    client = TestClient(app)
    run_id = _preview(client, tmp_path)

    started = client.post(f"/runs/{run_id}/start", follow_redirects=False)
    assert started.status_code == 303
    app.state.job_controller.wait(run_id, timeout=10)
    assert (
        app.state.run_store.get_run(run_id).status
        is RunStatus.AWAITING_PARSE_RULES_REVIEW
    )

    review_page = client.get(f"/runs/{run_id}/parse-rules")
    assert review_page.status_code == 200
    assert "Dictionary name" not in review_page.text

    approved = client.post(
        f"/runs/{run_id}/parse-rules/approve",
        follow_redirects=False,
    )
    assert approved.status_code == 303
    app.state.job_controller.wait(run_id, timeout=10)

    assert app.state.run_store.get_run(run_id).status is RunStatus.COMPLETED
    assert (tmp_path / "output/stage-2/page_1/page_1_mdf.txt").is_file()


def test_missing_key_moves_prepared_run_to_credentials_required(
    tmp_path: Path,
) -> None:
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=CredentialVault(environ={}),
        offline_inference=True,
    )

    client = TestClient(app)
    run_id = _preview(client, tmp_path)

    response = client.post(f"/runs/{run_id}/start")

    assert response.status_code == 409
    assert "API credential required" in response.text
    assert app.state.run_store.get_run(run_id).status is RunStatus.CREDENTIALS_REQUIRED

def test_credentials_required_page_has_one_provider_recovery_card(
    tmp_path: Path,
) -> None:
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=CredentialVault(environ={}),
        offline_inference=True,
    )
    client = TestClient(app)
    run_id = _preview(client, tmp_path)

    response = client.post(f"/runs/{run_id}/start")

    assert response.status_code == 409
    assert response.text.count('data-credential-card') == 1
    assert f'data-provider="anthropic"' in response.text
    assert 'data-save-key' in response.text
    assert f'data-continue-action="/runs/{run_id}/start"' in response.text
    assert f'href="/runs/{run_id}/review"' in response.text
    assert "Save and continue" in response.text
    assert "/static/app.js?v=dashboard-ui-8" in response.text
    assert "api_key" not in response.text


def test_prepared_review_snapshot_contains_no_credential(tmp_path: Path) -> None:
    vault = CredentialVault(environ={})
    vault.set_temporary(Provider.ANTHROPIC, "sk-ant-must-not-persist")
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=vault,
        offline_inference=True,
    )
    client = TestClient(app)
    run_id = _preview(client, tmp_path)

    config_text = app.state.job_controller.config_path(run_id).read_text(
        encoding="utf-8"
    )

    assert "sk-ant-must-not-persist" not in config_text
    assert "api_key" not in config_text.lower()


def test_live_log_is_managed_bounded_and_redacts_provider_key(tmp_path: Path) -> None:
    vault = CredentialVault(environ={})
    vault.set_temporary(Provider.ANTHROPIC, "sk-ant-live-log-secret")
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=vault,
        offline_inference=True,
    )
    client = TestClient(app)
    run_id = _preview(client, tmp_path)
    app.state.run_store.transition(run_id, RunStatus.QUEUED)
    app.state.run_store.transition(run_id, RunStatus.RUNNING_STAGE1)
    log_path = app.state.job_controller.log_path(run_id)
    log_path.write_text(
        ("old output\n" * 70_000) + "request sk-ant-live-log-secret failed\n",
        encoding="utf-8",
    )

    response = client.get(f"/runs/{run_id}/logs")

    assert response.status_code == 200
    assert "Live Logs" in response.text
    assert "[REDACTED]" in response.text
    assert "sk-ant-live-log-secret" not in response.text
    assert "Older log output was truncated" in response.text
    assert response.text.count('<pre data-log-console') == 1
    for marker in (
        'data-stream-status',
        'data-live-toggle',
        'data-log-copy',
        "Pause",
        "Resume",
        "Copy visible text",
    ):
        assert marker in response.text
    assert 'meta name="mudidi-events"' in response.text


@pytest.mark.parametrize(
    "terminal_event_type",
    ["run.completed", "run.failed", "run.cancelled"],
)
def test_logs_use_terminal_event_as_inactive_state_for_stale_run(
    tmp_path: Path,
    terminal_event_type: str,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    run_id = _preview(client, tmp_path)
    store = app.state.run_store
    store.transition(run_id, RunStatus.QUEUED)
    store.transition(run_id, RunStatus.RUNNING_STAGE1)
    terminal_event = {
        "version": 1,
        "type": terminal_event_type,
        "run_id": run_id,
        "sequence": 1,
        "occurred_at": "2026-09-06T00:00:00+00:00",
        "stage": "stage1",
    }
    if terminal_event_type == "run.failed":
        terminal_event["message"] = "persisted terminal failure"
    store.append_event(run_id, terminal_event)

    response = client.get(f"/runs/{run_id}/logs")

    assert response.status_code == 200
    assert 'meta name="mudidi-events"' not in response.text
    assert 'data-stream-status role="status">Idle</span>' in response.text

    stream = client.get(f"/runs/{run_id}/events?after=1")
    assert stream.status_code == 200
    assert stream.text == ""

def test_failed_run_surfaces_error_in_overview_and_logs(tmp_path: Path) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    run_id = _preview(client, tmp_path)
    app.state.run_store.transition(run_id, RunStatus.QUEUED)
    app.state.job_controller.start_fake(
        run_id,
        page_count=1,
        delay_seconds=0,
        fail=True,
    )
    app.state.job_controller.wait(run_id, timeout=10)

    overview = client.get(f"/runs/{run_id}")
    logs = client.get(f"/runs/{run_id}/logs")

    assert overview.status_code == 200
    assert "Failure details" in overview.text
    assert "Offline worker failure requested" in overview.text
    assert 'aria-disabled="true">MDF parsing guide</span>' in overview.text
    assert logs.status_code == 200
    assert "Run failure" in logs.text
    assert "Offline worker failure requested" in logs.text


def test_saved_preset_loads_into_editable_new_run_and_reuses_inputs(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    run_id = _preview(client, tmp_path)
    run_bundle = app.state.inputs.bundle(run_id)
    guide = run_bundle / "mdf_guide" / "saved-guide.json"
    guide.parent.mkdir()
    guide.write_text(
        json.dumps(
            {"markers": [{"marker": "lx", "description": "Headword"}]}
        ),
        encoding="utf-8",
    )
    manual = run_bundle / "mdf_manual" / "saved-manual.pdf"
    manual.parent.mkdir()
    manual.write_bytes(_pdf_bytes())
    run_config = app.state.job_controller.load_inference_config(run_id)
    run_config = run_config.model_copy(
        update={
            "input": run_config.input.model_copy(update={"toolbox_pdf": manual}),
            "pipeline": run_config.pipeline.model_copy(
                update={"parse_rules_file": guide}
            ),
        }
    )
    app.state.job_controller.config_path(run_id).write_text(
        run_config.model_dump_json(indent=2),
        encoding="utf-8",
    )

    saved = client.post(
        f"/runs/{run_id}/presets",
        data={"name": "My verified setup"},
        follow_redirects=False,
    )

    assert saved.status_code == 303
    preset = app.state.run_store.list_presets()[0]
    assert preset.config.input.pages is not None
    assert preset.config.input.pages.resolve().is_relative_to(
        (tmp_path / "app-data" / "presets" / preset.preset_id / "inputs").resolve()
    )
    page = client.get("/presets")
    assert page.status_code == 200
    assert "My verified setup" in page.text
    assert f'href="/?preset={preset.preset_id}"' in page.text
    assert "Use preset" in page.text
    assert f'aria-label="Use preset My verified setup"' in page.text
    assert "Updated" in page.text
    assert "Provider" in page.text
    assert "Pipeline" in page.text
    assert "Primary model" in page.text
    assert "Agentic" in page.text
    assert "Remove" in page.text
    assert f'aria-label="Remove preset My verified setup"' in page.text
    assert 'data-confirm-preset-delete="My verified setup"' in page.text
    assert 'action="/presets/' + preset.preset_id + '/delete"' in page.text

    home = client.get("/")
    assert 'name="preset"' in home.text
    assert "My verified setup" in home.text
    assert "Load preset" in home.text

    loaded = client.get(f"/?preset={preset.preset_id}")
    assert loaded.status_code == 200
    assert f'name="preset_id" value="{preset.preset_id}"' in loaded.text
    assert 'id="preset-form-state"' in loaded.text
    assert str(tmp_path / "output") in loaded.text
    assert "Loaded preset: My verified setup" in loaded.text
    assert 'id="dictionary-pdf"' in loaded.text
    assert re.search(r'id="dictionary-pdf"[^>]*disabled', loaded.text) is None
    assert re.search(r'name="existing_mdf_guide_file"[^>]*disabled', loaded.text) is None
    assert "dictionary.pdf" in loaded.text
    assert "saved-guide.json" in loaded.text
    assert "saved-manual.pdf" in loaded.text
    assert f'/presets/{preset.preset_id}/files/pages/0' in loaded.text
    assert f'/presets/{preset.preset_id}/files/mdf-guide' in loaded.text
    assert f'/presets/{preset.preset_id}/files/mdf-manual' in loaded.text

    saved_page = client.get(f"/presets/{preset.preset_id}/files/pages/0")
    saved_guide = client.get(f"/presets/{preset.preset_id}/files/mdf-guide")
    saved_manual = client.get(f"/presets/{preset.preset_id}/files/mdf-manual")
    assert saved_page.content.startswith(b"%PDF-")
    assert saved_guide.json()["markers"][0]["marker"] == "lx"
    assert saved_manual.content.startswith(b"%PDF-")

    # Presets own their input assets and remain valid after the source run is removed.
    shutil.rmtree(tmp_path / "app-data" / "runs" / run_id / "inputs")

    edited_output = tmp_path / "edited-output"
    review = client.post(
        "/runs/preview",
        data={
            "preset_id": preset.preset_id,
            "output_directory": str(edited_output),
            "output_policy": "resume",
            "pipeline": "complete",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "agentic": "false",
            "parse_rules_pages": "1",
            "dictionary_pages": "1",
        },
    )

    assert review.status_code == 200
    assert "Review your run" in review.text
    assert str(edited_output) in review.text
    assert len(app.state.run_store.list_runs()) == 2
    cloned = next(run for run in app.state.run_store.list_runs() if run.run_id != run_id)
    cloned_config = app.state.job_controller.load_inference_config(cloned.run_id)
    assert cloned_config.input.pages is not None
    assert cloned_config.input.pages.resolve().is_relative_to(
        (tmp_path / "app-data" / "runs" / cloned.run_id / "inputs").resolve()
    )

    replacement = client.post(
        "/runs/preview",
        data={
            "preset_id": preset.preset_id,
            "output_directory": str(tmp_path / "replacement-output"),
            "output_policy": "resume",
            "pipeline": "complete",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "agentic": "false",
            "mdf_manual_source": "upload",
            "dictionary_pages": "1",
        },
        files=[
            (
                "dictionary_pdf",
                ("replacement.pdf", _pdf_bytes(), "application/pdf"),
            ),
            (
                "existing_mdf_guide_file",
                (
                    "replacement-guide.json",
                    json.dumps(
                        {"markers": [{"marker": "ge", "description": "Gloss"}]}
                    ),
                    "application/json",
                ),
            ),
            ("custom_mdf_manual", ("replacement.pdf", _pdf_bytes(), "application/pdf")),
        ],
    )

    assert replacement.status_code == 200
    replacement_run = app.state.run_store.list_runs()[0]
    replacement_config = app.state.job_controller.load_inference_config(
        replacement_run.run_id
    )
    assert replacement_config.input.pages.name == "replacement.pdf"
    assert replacement_config.input.pages.is_file()
    assert replacement_config.pipeline.parse_rules_file.name == "replacement-guide.json"
    assert replacement_config.input.toolbox_pdf.name == "replacement.pdf"

    overwritten = client.post(
        f"/runs/{replacement_run.run_id}/presets",
        data={"name": "My verified setup"},
        follow_redirects=False,
    )
    assert overwritten.status_code == 303
    saved_presets = app.state.run_store.list_presets()
    assert len(saved_presets) == 1
    assert saved_presets[0].preset_id != preset.preset_id
    assert not (tmp_path / "app-data" / "presets" / preset.preset_id).exists()
    assert client.get(f"/presets/{preset.preset_id}/files/pages/0").status_code == 404


def test_preset_remove_deletes_metadata_bundle_and_returns_to_presets(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data")
    client = TestClient(app)
    pages = tmp_path / "pages"
    pages.mkdir()
    preset = app.state.run_store.create_preset(
        "preset-remove",
        name="Remove me",
        provider="offline",
        config=InferenceConfig.model_validate(
            {
                "input": {"pages": pages},
                "output": {"directory": tmp_path / "output"},
            }
        ),
    )
    bundle = app.state.inputs.presets_root / preset.preset_id / "inputs"
    bundle.mkdir(parents=True)
    (bundle / "source.pdf").write_bytes(_pdf_bytes())

    response = client.post(
        f"/presets/{preset.preset_id}/delete",
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"] == "/presets"
    assert not bundle.parent.exists()
    assert client.get("/presets").status_code == 200
    assert "Remove me" not in client.get("/presets").text
    assert client.post("/presets/unknown/delete").status_code == 404
    empty = client.get("/presets")
    assert "No presets yet" in empty.text
    assert 'href="/">' in empty.text
    assert "Create a new run" in empty.text


def test_interrupted_stage1_run_can_resume_from_run_detail(tmp_path: Path) -> None:
    vault = CredentialVault(environ={})
    vault.set_temporary(Provider.ANTHROPIC, "sk-ant-resume")
    app = create_app(
        data_dir=tmp_path / "app-data",
        credential_vault=vault,
        offline_inference=True,
    )
    client = TestClient(app)
    run_id = _preview(client, tmp_path)
    app.state.run_store.transition(run_id, RunStatus.QUEUED)
    app.state.run_store.transition(run_id, RunStatus.RUNNING_STAGE1)
    app.state.run_store.interrupt(run_id)

    detail = client.get(f"/runs/{run_id}")
    assert "Resume run" in detail.text

    response = client.post(f"/runs/{run_id}/resume", follow_redirects=False)

    assert response.status_code == 303
    app.state.job_controller.wait(run_id, timeout=10)
    assert (
        app.state.run_store.get_run(run_id).status
        is RunStatus.AWAITING_PARSE_RULES_REVIEW
    )


def test_preview_materializes_stage1_selected_pdf_instruction(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    response = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "output"),
            "pipeline": "transcription",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "file",
            "stage1_instruction_pdf_pages": "2-3",
        },
        files={
            "dictionary_pdf": ("dictionary.pdf", _pdf_bytes(), "application/pdf"),
            "stage1_instruction_file": (
                "stage1-reference.pdf",
                _pdf_bytes(3),
                "application/pdf",
            ),
        },
    )

    assert response.status_code == 200
    run_match = re.search(r'action="/runs/([^/]+)/start"', response.text)
    assert run_match is not None
    run_id = run_match.group(1)
    config = app.state.job_controller.load_inference_config(run_id)
    assert config.pipeline.stage1_guides is not None
    assert config.pipeline.stage1_guides.parent.name == "stage1"
    assert config.pipeline.stage1_guides_pages == "2-3"
    metadata = json.loads(
        (config.pipeline.stage1_guides.parent / "metadata.json").read_text(
            encoding="utf-8"
        )
    )
    assert metadata["source_mode"] == "file"
    assert metadata["kind"] == "pdf"
    assert metadata["pdf_page_count"] == 3
    assert metadata["selected_pages"] == [2, 3]


def test_preview_materializes_stage2_markdown_for_pass1_only(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    guide_text = "# Parse-guide discovery\\nKeep field roles explicit."
    response = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "output"),
            "pipeline": "structure",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage2_instruction_source": "file",
            "stage2_instruction_scope": "pass1",
        },
        files={
            "dictionary_pdf": ("dictionary.pdf", _pdf_bytes(), "application/pdf"),
            "stage2_instruction_file": (
                "stage2-guide.md",
                guide_text.encode(),
                "text/markdown",
            ),
        },
    )

    assert response.status_code == 200
    run_match = re.search(r'action="/runs/([^/]+)/start"', response.text)
    assert run_match is not None
    run_id = run_match.group(1)
    config = app.state.job_controller.load_inference_config(run_id)
    assert config.pipeline.stage2_guides is not None
    assert config.pipeline.stage2_guides.parent.name == "stage2"
    assert config.pipeline.stage2_guides.read_text(encoding="utf-8") == guide_text
    assert config.pipeline.stage2_guides_scope == "pass1"
    metadata = json.loads(
        (config.pipeline.stage2_guides.parent / "metadata.json").read_text(
            encoding="utf-8"
        )
    )
    assert metadata["source_mode"] == "file"
    assert metadata["kind"] == "text"
    assert metadata["selected_pages"] == []
    assert metadata["stage2_scope"] == "pass1"


def test_preview_materializes_instruction_uploads_and_review_metadata(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    stage1_text = "PRIVATE stage one instruction"
    stage2_pdf = _pdf_bytes(3)

    response = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "output"),
            "pipeline": "complete",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "file",
            "stage1_instruction_pdf_pages": "",
            "stage2_instruction_source": "file",
            "stage2_instruction_pdf_pages": "2-3,2",
            "stage2_instruction_scope": "pass1",
        },
        files=[
            ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
            (
                "stage1_instruction_file",
                ("notes.txt", stage1_text.encode(), "text/plain"),
            ),
            (
                "stage2_instruction_file",
                ("reference.pdf", stage2_pdf, "application/pdf"),
            ),
        ],
    )

    assert response.status_code == 200
    assert stage1_text not in response.text
    assert "notes.txt" in response.text
    assert "reference.pdf" in response.text
    assert "Pass 1 only" in response.text
    run_id = re.search(r'action="/runs/([^/]+)/start"', response.text)
    assert run_id is not None
    config = app.state.job_controller.load_inference_config(run_id.group(1))
    assert config.pipeline.stage1_guides is not None
    assert config.pipeline.stage1_guides.parent.name == "stage1"
    assert config.pipeline.stage2_guides is not None
    assert config.pipeline.stage2_guides.parent.name == "stage2"
    assert config.pipeline.stage2_guides_pages == "2-3"
    assert config.pipeline.stage2_guides_scope == "pass1"


@pytest.mark.parametrize(
    ("source", "text", "filename"),
    [
        ("typed", "typed text", "forbidden.txt"),
        ("file", "typed text", "guide.txt"),
    ],
)
def test_preview_rejects_instruction_source_mixing(
    tmp_path: Path,
    source: str,
    text: str,
    filename: str,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    data = {
        "output_directory": str(tmp_path / "output"),
        "pipeline": "complete",
        "dictionary_pages": "1",
        "provider": "anthropic",
        "model": "anthropic/claude-sonnet-5",
        "reasoning": "low",
        "stage1_instruction_source": source,
        "stage1_additional_instructions": text,
    }
    files = {"dictionary_pdf": ("dictionary.pdf", _pdf_bytes(), "application/pdf")}
    if source == "typed":
        files["stage1_instruction_file"] = (filename, b"uploaded", "text/plain")

    response = client.post("/runs/preview", data=data, files=files)

    assert response.status_code == 422
    assert "Remove the uploaded file" in response.text or (
        "Clear typed instructions" in response.text
    )
    assert not app.state.run_store.list_runs()


def test_preview_rejects_forged_stage2_values_when_pipeline_is_inactive(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)

    response = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "output"),
            "pipeline": "transcription",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage2_instruction_source": "file",
            "stage2_instruction_scope": "pass1",
            "stage2_instruction_pdf_pages": "1",
        },
        files=[
            ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
            ("stage2_instruction_file", ("guide.txt", b"forbidden", "text/plain")),
        ],
    )

    assert response.status_code == 422
    assert not app.state.run_store.list_runs()


def test_instruction_preset_keep_replace_and_clear_preserves_sidecars(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    response = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "source-output"),
            "pipeline": "complete",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "file",
            "stage2_instruction_source": "file",
            "stage2_instruction_scope": "both",
        },
        files=[
            ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
            ("stage1_instruction_file", ("stage1.txt", b"one", "text/plain")),
            ("stage2_instruction_file", ("stage2.md", b"two", "text/markdown")),
        ],
    )
    assert response.status_code == 200
    source_run_id = re.search(r'action="/runs/([^/]+)/start"', response.text)
    assert source_run_id is not None
    source_run_id = source_run_id.group(1)

    saved = client.post(
        f"/runs/{source_run_id}/presets",
        data={"name": "Instruction preset"},
        follow_redirects=False,
    )
    assert saved.status_code == 303
    preset = app.state.run_store.list_presets()[0]
    preset_root = app.state.inputs.presets_root / preset.preset_id / "inputs"
    stage1_source = preset_root / "instructions" / "stage1" / "stage1.txt"
    stage2_source = preset_root / "instructions" / "stage2" / "stage2.md"
    assert (stage1_source.parent / "metadata.json").is_file()
    assert (stage2_source.parent / "metadata.json").is_file()
    loaded = client.get(f"/?preset={preset.preset_id}")
    assert loaded.status_code == 200
    assert "stage1_instruction_source" in loaded.text
    assert "stage2_instruction_source" in loaded.text
    links = _preset_asset_links(
        preset,
        presets_root=app.state.inputs.presets_root,
    )
    assert links["stage1_instruction"]["name"] == "stage1.txt"
    assert links["stage2_instruction"]["name"] == "stage2.md"

    kept = client.post(
        "/runs/preview",
        data={
            "preset_id": preset.preset_id,
            "output_directory": str(tmp_path / "kept-output"),
            "pipeline": "complete",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "file",
            "stage1_instruction_keep_existing": "true",
            "stage2_instruction_source": "file",
            "stage2_instruction_keep_existing": "true",
            "stage2_instruction_scope": "both",
        },
    )
    assert kept.status_code == 200
    kept_run_id = re.search(r'action="/runs/([^/]+)/start"', kept.text)
    assert kept_run_id is not None
    kept_config = app.state.job_controller.load_inference_config(kept_run_id.group(1))
    assert kept_config.pipeline.stage1_guides is not None
    assert kept_config.pipeline.stage1_guides.read_text(encoding="utf-8") == "one"
    assert kept_config.pipeline.stage2_guides is not None
    assert kept_config.pipeline.stage2_guides.read_text(encoding="utf-8") == "two"
    assert (kept_config.pipeline.stage1_guides.parent / "metadata.json").is_file()

    replaced = client.post(
        "/runs/preview",
        data={
            "preset_id": preset.preset_id,
            "output_directory": str(tmp_path / "replaced-output"),
            "pipeline": "complete",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "file",
            "stage2_instruction_source": "file",
            "stage2_instruction_keep_existing": "true",
            "stage2_instruction_scope": "both",
        },
        files={
            "stage1_instruction_file": (
                "replacement.txt",
                b"replacement",
                "text/plain",
            )
        },
    )
    assert replaced.status_code == 200
    replaced_run_id = re.search(r'action="/runs/([^/]+)/start"', replaced.text)
    assert replaced_run_id is not None
    replaced_config = app.state.job_controller.load_inference_config(
        replaced_run_id.group(1)
    )
    assert replaced_config.pipeline.stage1_guides is not None
    assert replaced_config.pipeline.stage1_guides.name == "replacement.txt"
    assert replaced_config.pipeline.stage1_guides.read_text(encoding="utf-8") == "replacement"

    cleared = client.post(
        "/runs/preview",
        data={
            "preset_id": preset.preset_id,
            "output_directory": str(tmp_path / "cleared-output"),
            "pipeline": "complete",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "typed",
            "stage1_additional_instructions": "",
            "stage2_instruction_source": "file",
            "stage2_instruction_keep_existing": "true",
            "stage2_instruction_scope": "both",
        },
    )
    assert cleared.status_code == 200
    cleared_run_id = re.search(r'action="/runs/([^/]+)/start"', cleared.text)
    assert cleared_run_id is not None
    cleared_config = app.state.job_controller.load_inference_config(
        cleared_run_id.group(1)
    )
    assert cleared_config.pipeline.stage1_guides is None
    assert cleared_config.pipeline.stage2_guides is not None
    assert not (
        app.state.inputs.bundle(cleared_run_id.group(1))
        / "instructions"
        / "stage1"
    ).exists()


def test_sidecar_free_legacy_docx_preset_restores_extracted_text(
    tmp_path: Path,
) -> None:
    legacy = tmp_path / "legacy-guide.docx"
    document = Document()
    document.add_paragraph("Restore this extracted instruction.")
    document.save(str(legacy))

    assert _read_preset_text(legacy) == "Restore this extracted instruction."


@pytest.mark.parametrize(
    ("data", "files", "expected_field"),
    [
        (
            [
                ("stage1_instruction_source", "typed"),
                ("stage1_instruction_source", "file"),
            ],
            {},
            "Stage1 Instruction File",
        ),
        (
            [
                ("stage1_additional_instructions", "first"),
                ("stage1_additional_instructions", "second"),
            ],
            {},
            "Stage1 Instruction File",
        ),
        (
            [
                ("stage1_instruction_source", "file"),
                ("stage1_instruction_pdf_pages", "1"),
                ("stage1_instruction_pdf_pages", "2"),
            ],
            {
                "stage1_instruction_file": (
                    "guide.pdf",
                    _pdf_bytes(2),
                    "application/pdf",
                )
            },
            "Stage1 Instruction Pdf Pages",
        ),
    ],
)
def test_preview_rejects_ambiguous_instruction_scalars(
    tmp_path: Path,
    data: list[tuple[str, str]],
    files: dict[str, tuple[str, bytes, str]],
    expected_field: str,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    multipart: list[tuple[str, object]] = [
        ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
        ("output_directory", (None, str(tmp_path / "output"))),
        ("pipeline", (None, "complete")),
        ("dictionary_pages", (None, "1")),
        ("provider", (None, "anthropic")),
        ("model", (None, "anthropic/claude-sonnet-5")),
        ("reasoning", (None, "low")),
    ]
    multipart.extend((field, (None, value)) for field, value in data)
    multipart.extend(files.items())
    response = client.post("/runs/preview", files=multipart)

    assert response.status_code == 422
    assert f"<strong>{expected_field}:</strong>" in response.text
    assert not app.state.run_store.list_runs()


@pytest.mark.parametrize(
    ("pipeline", "source", "page_spec", "content", "expected_field"),
    [
        (
            "transcription",
            "file",
            None,
            b"forbidden",
            "Stage2 Instruction File",
        ),
        (
            "complete",
            "file",
            None,
            b"%PDF-",
            "Stage1 Instruction File",
        ),
        (
            "complete",
            "file",
            "3",
            _pdf_bytes(2),
            "Stage1 Instruction Pdf Pages",
        ),
        (
            "complete",
            "file",
            "1",
            b"typed guide",
            "Stage1 Instruction Pdf Pages",
        ),
    ],
)
def test_preview_instruction_failures_use_safe_error_key(
    tmp_path: Path,
    pipeline: str,
    source: str,
    page_spec: str | None,
    content: bytes,
    expected_field: str,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    data: list[tuple[str, str]] = [
        ("output_directory", str(tmp_path / "output")),
        ("pipeline", pipeline),
        ("dictionary_pages", "1"),
        ("provider", "anthropic"),
        ("model", "anthropic/claude-sonnet-5"),
        ("reasoning", "low"),
        ("stage1_instruction_source", source),
        ("stage2_instruction_source", source if pipeline == "transcription" else "typed"),
    ]
    if page_spec is not None:
        data.append(("stage1_instruction_pdf_pages", page_spec))
    stage = "stage2" if pipeline == "transcription" else "stage1"
    file_field = f"{stage}_instruction_file"
    filename = "guide.pdf" if content.startswith(b"%PDF-") else "guide.txt"
    if stage == "stage2":
        data = [
            item
            for item in data
            if item[0] != "stage1_instruction_source"
        ]
    multipart: list[tuple[str, object]] = [
        ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
    ]
    multipart.extend((field, (None, value)) for field, value in data)
    multipart.append(
        (
            file_field,
            (
                filename,
                content,
                "application/pdf"
                if filename.endswith(".pdf")
                else "text/plain",
            ),
        )
    )
    response = client.post("/runs/preview", files=multipart)

    assert response.status_code == 422
    assert f"<strong>{expected_field}:</strong>" in response.text
    assert not app.state.run_store.list_runs()


def test_preview_normalizes_whitespace_padded_instruction_source_and_scope(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    multipart: list[tuple[str, object]] = [
        ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
        ("output_directory", (None, str(tmp_path / "output"))),
        ("pipeline", (None, "complete")),
        ("dictionary_pages", (None, "1")),
        ("provider", (None, "anthropic")),
        ("model", (None, "anthropic/claude-sonnet-5")),
        ("reasoning", (None, "low")),
        ("stage1_instruction_source", (None, " file ")),
        ("stage2_instruction_source", (None, " file ")),
        ("stage2_instruction_scope", (None, " pass1 ")),
        ("stage1_instruction_file", ("stage1.txt", b"stage one", "text/plain")),
        (
            "stage2_instruction_file",
            ("stage2.txt", b"stage two", "text/plain"),
        ),
    ]

    response = client.post("/runs/preview", files=multipart)

    assert response.status_code == 200
    run_id = re.search(r'action="/runs/([^/]+)/start"', response.text)
    assert run_id is not None
    config = app.state.job_controller.load_inference_config(run_id.group(1))
    assert config.pipeline.stage1_guides is not None
    assert config.pipeline.stage2_guides is not None
    assert config.pipeline.stage2_guides_scope == "pass1"


def test_preview_whitespace_padded_non_pdf_page_selection_uses_pages_error(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    multipart: list[tuple[str, object]] = [
        ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
        ("output_directory", (None, str(tmp_path / "output"))),
        ("pipeline", (None, "transcription")),
        ("dictionary_pages", (None, "1")),
        ("provider", (None, "anthropic")),
        ("model", (None, "anthropic/claude-sonnet-5")),
        ("reasoning", (None, "low")),
        ("stage1_instruction_source", (None, " file ")),
        ("stage1_instruction_pdf_pages", (None, " 1 ")),
        ("stage1_instruction_file", ("stage1.txt", b"stage one", "text/plain")),
    ]

    response = client.post("/runs/preview", files=multipart)

    assert response.status_code == 422
    assert "<strong>Stage1 Instruction Pdf Pages:</strong>" in response.text
    assert "Stage1 Instruction Source" not in response.text
    assert not app.state.run_store.list_runs()


def test_instruction_validation_preserves_safe_non_file_form_state(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    output = tmp_path / "safe-output"
    multipart: list[tuple[str, object]] = [
        ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
        ("output_directory", (None, str(output))),
        ("pipeline", (None, "transcription")),
        ("dictionary_pages", (None, "1")),
        ("provider", (None, "anthropic")),
        ("model", (None, "anthropic/claude-sonnet-5")),
        ("reasoning", (None, "low")),
        ("stage1_instruction_source", (None, "file")),
        (
            "stage1_instruction_file",
            ("guide.txt", b" \n", "text/plain"),
        ),
    ]

    response = client.post("/runs/preview", files=multipart)

    assert response.status_code == 422
    assert "<strong>Stage1 Instruction File:</strong>" in response.text
    assert "Stage1 Instruction Source" not in response.text
    assert "Stage1 Instruction Pdf Pages" not in response.text
    assert str(output) not in response.text
    assert 'value="~/Documents/MUDIDI-runs"' in response.text
    assert "private instruction" not in response.text
    assert not app.state.run_store.list_runs()


def test_uploaded_instruction_review_matches_recovered_review_metadata(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    response = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "output"),
            "pipeline": "complete",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "file",
            "stage2_instruction_source": "file",
            "stage2_instruction_pdf_pages": "2",
            "stage2_instruction_scope": "pass2",
        },
        files=[
            ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
            ("stage1_instruction_file", ("stage1.txt", b"private one", "text/plain")),
            (
                "stage2_instruction_file",
                ("stage2.pdf", _pdf_bytes(3), "application/pdf"),
            ),
        ],
    )

    assert response.status_code == 200
    run_match = re.search(r'action="/runs/([^/]+)/start"', response.text)
    assert run_match is not None
    recovered = client.get(f"/runs/{run_match.group(1)}/review")
    assert recovered.status_code == 200

    def rows(html: str) -> dict[str, str]:
        return {
            stage: re.search(
                rf"<dt>{stage} Instructions</dt>\s*<dd>(.*?)</dd>",
                html,
                re.S,
            ).group(1)
            for stage in ("Stage 1", "Stage 2")
        }

    assert rows(response.text) == rows(recovered.text)
    assert "private one" not in response.text
    assert "private one" not in recovered.text


def test_pdf_preset_restores_pages_and_explicit_blank_keep_preserves_all_pages(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)

    selected = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "selected-output"),
            "pipeline": "structure",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage2_instruction_source": "file",
            "stage2_instruction_pdf_pages": "2",
            "stage2_instruction_scope": "pass2",
        },
        files=[
            ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
            (
                "stage2_instruction_file",
                ("selected.pdf", _pdf_bytes(3), "application/pdf"),
            ),
        ],
    )
    assert selected.status_code == 200
    selected_run = re.search(r'action="/runs/([^/]+)/start"', selected.text)
    assert selected_run is not None
    saved = client.post(
        f"/runs/{selected_run.group(1)}/presets",
        data={"name": "Selected PDF preset"},
        follow_redirects=False,
    )
    assert saved.status_code == 303
    selected_preset = app.state.run_store.list_presets()[0]

    loaded = client.get(f"/?preset={selected_preset.preset_id}")
    assert loaded.status_code == 200
    state_match = re.search(
        r'<script id="preset-form-state" type="application/json">(.*?)</script>',
        loaded.text,
        re.S,
    )
    assert state_match is not None
    state = json.loads(state_match.group(1))
    assert state["stage2_instruction_pdf_pages"] == ["2"]
    assert state["stage2_instruction_scope"] == ["pass2"]

    kept = client.post(
        "/runs/preview",
        data={
            "preset_id": selected_preset.preset_id,
            "output_directory": str(tmp_path / "selected-kept-output"),
            "pipeline": "structure",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage2_instruction_source": "file",
            "stage2_instruction_pdf_pages": "2",
            "stage2_instruction_keep_existing": "true",
            "stage2_instruction_scope": "pass2",
        },
    )
    assert kept.status_code == 200
    kept_run = re.search(r'action="/runs/([^/]+)/start"', kept.text)
    assert kept_run is not None
    kept_config = app.state.job_controller.load_inference_config(kept_run.group(1))
    assert kept_config.pipeline.stage2_guides_pages == "2"
    assert kept_config.pipeline.stage2_guides_scope == "pass2"

    all_pages = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "all-pages-output"),
            "pipeline": "structure",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage2_instruction_source": "file",
        },
        files=[
            ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
            (
                "stage2_instruction_file",
                ("all.pdf", _pdf_bytes(3), "application/pdf"),
            ),
        ],
    )
    assert all_pages.status_code == 200
    all_run = re.search(r'action="/runs/([^/]+)/start"', all_pages.text)
    assert all_run is not None
    all_saved = client.post(
        f"/runs/{all_run.group(1)}/presets",
        data={"name": "All PDF preset"},
        follow_redirects=False,
    )
    assert all_saved.status_code == 303
    all_preset = next(
        preset
        for preset in app.state.run_store.list_presets()
        if preset.name == "All PDF preset"
    )

    blank_kept = client.post(
        "/runs/preview",
        data={
            "preset_id": all_preset.preset_id,
            "output_directory": str(tmp_path / "all-pages-kept-output"),
            "pipeline": "structure",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage2_instruction_source": "file",
            "stage2_instruction_pdf_pages": " ",
            "stage2_instruction_keep_existing": "true",
            "stage2_instruction_scope": "both",
        },
    )
    assert blank_kept.status_code == 200
    blank_run = re.search(r'action="/runs/([^/]+)/start"', blank_kept.text)
    assert blank_run is not None
    blank_config = app.state.job_controller.load_inference_config(blank_run.group(1))
    assert blank_config.pipeline.stage2_guides_pages is None
    assert blank_config.pipeline.stage2_guides_scope == "both"
    metadata = json.loads(
        (
            blank_config.pipeline.stage2_guides.parent / "metadata.json"
        ).read_text(encoding="utf-8")
    )
    assert metadata["selected_pages"] == [1, 2, 3]


def test_typed_blank_deletes_managed_and_legacy_instruction_paths(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    managed = client.post(
        "/runs/preview",
        data={
            "output_directory": str(tmp_path / "managed-output"),
            "pipeline": "transcription",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "file",
        },
        files=[
            ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
            (
                "stage1_instruction_file",
                ("managed.txt", b"managed", "text/plain"),
            ),
        ],
    )
    assert managed.status_code == 200
    managed_run = re.search(r'action="/runs/([^/]+)/start"', managed.text)
    assert managed_run is not None
    saved = client.post(
        f"/runs/{managed_run.group(1)}/presets",
        data={"name": "Managed delete preset"},
        follow_redirects=False,
    )
    assert saved.status_code == 303
    managed_preset = app.state.run_store.list_presets()[0]
    cleared = client.post(
        "/runs/preview",
        data={
            "preset_id": managed_preset.preset_id,
            "output_directory": str(tmp_path / "managed-cleared-output"),
            "pipeline": "transcription",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "typed",
            "stage1_additional_instructions": "",
        },
    )
    assert cleared.status_code == 200
    cleared_run = re.search(r'action="/runs/([^/]+)/start"', cleared.text)
    assert cleared_run is not None
    cleared_bundle = app.state.inputs.bundle(cleared_run.group(1))
    assert not (cleared_bundle / "instructions" / "stage1").exists()

    legacy_root = app.state.inputs.presets_root / "legacy-delete" / "inputs"
    legacy_pages = legacy_root / "pages.pdf"
    legacy_guide = legacy_root / "instructions" / "stage1.txt"
    legacy_guide.parent.mkdir(parents=True)
    legacy_pages.parent.mkdir(parents=True, exist_ok=True)
    legacy_pages.write_bytes(_pdf_bytes())
    legacy_guide.write_text("legacy", encoding="utf-8")
    legacy_config = InferenceConfig.model_validate(
        {
            "input": {"pages": legacy_pages},
            "output": {"directory": tmp_path / "legacy-output"},
            "pipeline": {"stage": "1", "stage1_guides": legacy_guide},
        }
    )
    app.state.run_store.create_preset(
        "legacy-delete",
        name="Legacy delete preset",
        provider="anthropic",
        config=legacy_config,
    )
    legacy_cleared = client.post(
        "/runs/preview",
        data={
            "preset_id": "legacy-delete",
            "output_directory": str(tmp_path / "legacy-cleared-output"),
            "pipeline": "transcription",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "typed",
            "stage1_additional_instructions": "",
        },
    )
    assert legacy_cleared.status_code == 200
    legacy_run = re.search(r'action="/runs/([^/]+)/start"', legacy_cleared.text)
    assert legacy_run is not None
    legacy_bundle = app.state.inputs.bundle(legacy_run.group(1))
    assert not (legacy_bundle / "instructions" / "stage1.txt").exists()


@pytest.mark.parametrize(
    "field",
    ["stage2_instruction_scope", "stage2_instruction_keep_existing"],
)
def test_preview_rejects_repeated_stage2_scope_and_keep_controls(
    tmp_path: Path,
    field: str,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    multipart: list[tuple[str, object]] = [
        ("dictionary_pdf", ("dictionary.pdf", _pdf_bytes(), "application/pdf")),
        ("output_directory", (None, str(tmp_path / "output"))),
        ("pipeline", (None, "structure")),
        ("dictionary_pages", (None, "1")),
        ("provider", (None, "anthropic")),
        ("model", (None, "anthropic/claude-sonnet-5")),
        ("reasoning", (None, "low")),
        ("stage2_instruction_source", (None, "file")),
        (field, (None, "pass1" if "scope" in field else "true")),
        (field, (None, "pass2" if "scope" in field else "false")),
    ]
    response = client.post("/runs/preview", files=multipart)

    assert response.status_code == 422
    assert "<strong>Stage2 Instruction File:</strong>" in response.text
    assert not app.state.run_store.list_runs()


def test_real_legacy_docx_preset_restores_typed_instruction_state(
    tmp_path: Path,
) -> None:
    app = create_app(data_dir=tmp_path / "app-data", offline_inference=True)
    client = TestClient(app)
    preset_root = app.state.inputs.presets_root / "legacy-docx" / "inputs"
    pages = preset_root / "pages.pdf"
    guide = preset_root / "instructions" / "stage1.docx"
    guide.parent.mkdir(parents=True)
    pages.write_bytes(_pdf_bytes())
    document = Document()
    document.add_paragraph("Restore this real DOCX instruction.")
    document.save(str(guide))
    config = InferenceConfig.model_validate(
        {
            "input": {"pages": pages},
            "output": {"directory": tmp_path / "legacy-docx-output"},
            "pipeline": {"stage": "1", "stage1_guides": guide},
        }
    )
    app.state.run_store.create_preset(
        "legacy-docx",
        name="Legacy DOCX preset",
        provider="anthropic",
        config=config,
    )

    loaded = client.get("/?preset=legacy-docx")
    assert loaded.status_code == 200
    state_match = re.search(
        r'<script id="preset-form-state" type="application/json">(.*?)</script>',
        loaded.text,
        re.S,
    )
    assert state_match is not None
    state = json.loads(state_match.group(1))
    assert state["stage1_instruction_source"] == ["typed"]
    assert state["stage1_additional_instructions"] == [
        "Restore this real DOCX instruction."
    ]

    restored = client.post(
        "/runs/preview",
        data={
            "preset_id": "legacy-docx",
            "output_directory": str(tmp_path / "restored-output"),
            "pipeline": "transcription",
            "dictionary_pages": "1",
            "provider": "anthropic",
            "model": "anthropic/claude-sonnet-5",
            "reasoning": "low",
            "stage1_instruction_source": "typed",
            "stage1_additional_instructions": "Restore this real DOCX instruction.",
        },
    )
    assert restored.status_code == 200
    run_match = re.search(r'action="/runs/([^/]+)/start"', restored.text)
    assert run_match is not None
    restored_config = app.state.job_controller.load_inference_config(run_match.group(1))
    assert restored_config.pipeline.stage1_guides is not None
    assert (
        restored_config.pipeline.stage1_guides.read_text(encoding="utf-8")
        == "Restore this real DOCX instruction."
    )
