from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pydantic import ValidationError

from mudidi.config.yaml_config import InferenceConfig, validate_config_paths


def _config(
    tmp_path: Path,
    *,
    stage: str = "1",
    strategy: str = "two_stage",
    stage1_guides: Path | None = None,
    stage2_guides: Path | None = None,
    stage2_scope: str = "both",
    vlm_model: str | None = None,
) -> InferenceConfig:
    (tmp_path / "pages").mkdir(exist_ok=True)
    settings: dict[str, object] = {
        "version": 1,
        "kind": "inference",
        "input": {"pages": tmp_path / "pages"},
        "output": {"directory": tmp_path / "output"},
        "pipeline": {
            "stage": stage,
            "strategy": strategy,
            "stage1_guides": stage1_guides,
            "stage2_guides": stage2_guides,
            "stage2_guides_scope": stage2_scope,
        },
    }
    if vlm_model is not None:
        settings["vlm"] = {"model": vlm_model}
    return InferenceConfig.model_validate(settings)


def test_stage1_guide_must_intersect_selected_stage(tmp_path: Path) -> None:
    guide = tmp_path / "stage1.txt"
    guide.write_text("stage one", encoding="utf-8")

    with pytest.raises(ValidationError, match="stage1_guides.*Stage 1"):
        _config(tmp_path, stage="2", stage1_guides=guide)


def test_stage2_guide_must_intersect_selected_stage(tmp_path: Path) -> None:
    guide = tmp_path / "stage2.txt"
    guide.write_text("stage two", encoding="utf-8")

    with pytest.raises(ValidationError, match="stage2_guides.*Stage 2"):
        _config(tmp_path, stage="1", stage2_guides=guide)


@pytest.mark.parametrize(
    ("stage", "scope"),
    [("2-pass-2", "pass1"), ("2-pass-1", "pass2")],
)
def test_stage2_scope_must_intersect_selected_pass(
    tmp_path: Path,
    stage: str,
    scope: str,
) -> None:
    guide = tmp_path / "stage2.txt"
    guide.write_text("stage two", encoding="utf-8")

    with pytest.raises(ValidationError, match="stage2_guides_scope.*selected"):
        _config(tmp_path, stage=stage, stage2_guides=guide, stage2_scope=scope)


def test_default_stage2_scope_without_guide_is_harmless(tmp_path: Path) -> None:
    config = _config(tmp_path, stage="1")

    assert config.pipeline.stage2_guides is None
    assert config.pipeline.stage2_guides_scope == "both"


@pytest.mark.parametrize(
    ("stage", "scope"),
    [
        ("1", "both"),
        ("all", "pass1"),
        ("2", "pass1"),
        ("2", "pass2"),
        ("2", "both"),
        ("2-pass-1", "pass1"),
        ("2-pass-1", "both"),
        ("2-pass-2", "pass2"),
        ("2-pass-2", "both"),
    ],
)
def test_configured_guides_accept_valid_stage_and_pass_intersections(
    tmp_path: Path,
    stage: str,
    scope: str,
) -> None:
    stage1 = tmp_path / "stage1.txt"
    stage1.write_text("stage one", encoding="utf-8")
    stage2 = tmp_path / "stage2.txt"
    stage2.write_text("stage two", encoding="utf-8")
    if stage == "1":
        config = _config(tmp_path, stage=stage, stage1_guides=stage1)
    else:
        config = _config(
            tmp_path,
            stage=stage,
            stage2_guides=stage2,
            stage2_scope=scope,
        )

    assert config.pipeline.stage == stage


@pytest.mark.parametrize("suffix", [".txt", ".md", ".docx"])
@pytest.mark.parametrize("text", [" \n", "x" * 20_000, "x" * 20_001])
def test_config_text_guide_validation_matches_runtime(
    tmp_path: Path,
    suffix: str,
    text: str,
) -> None:
    guide = tmp_path / f"guide{suffix}"
    if suffix == ".docx":
        document = Document()
        document.add_paragraph(text)
        document.save(str(guide))
    else:
        guide.write_text(text, encoding="utf-8")
    config = _config(tmp_path, stage1_guides=guide)

    if not text.strip():
        expected = "blank"
    elif len(text) > 20_000:
        expected = "too large"
    else:
        expected = None
    if expected is None:
        validate_config_paths(config)
    else:
        with pytest.raises(ValueError, match=expected):
            validate_config_paths(config)


@pytest.mark.parametrize("strategy", ["vlm_ocr", "mathpix_ocr"])
@pytest.mark.parametrize("suffix", [".txt", ".md", ".docx"])
def test_ocr_strategies_keep_non_pdf_guide_support(
    tmp_path: Path,
    strategy: str,
    suffix: str,
) -> None:
    guide = tmp_path / f"ocr-guide{suffix}"
    if suffix == ".docx":
        document = Document()
        document.add_paragraph("OCR guide")
        document.save(str(guide))
    else:
        guide.write_text("OCR guide", encoding="utf-8")
    config = _config(
        tmp_path,
        strategy=strategy,
        stage1_guides=guide,
        vlm_model="glm-ocr" if strategy == "vlm_ocr" else None,
    )

    validate_config_paths(config)
